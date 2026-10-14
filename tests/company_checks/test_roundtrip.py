"""Actual HTTP CRM -> parser -> persisted worker -> Range xlsx, mocked source data."""

from __future__ import annotations

import io
import threading
from http.server import ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from fastapi import FastAPI
from openpyxl import Workbook, load_workbook

from app.modules.lavok_parser import checks


@pytest.mark.asyncio
async def test_upload_progress_and_download(tmp_path, monkeypatch):
    parser_root = Path(__file__).resolve().parents[3] / "firm-parser"
    if not parser_root.is_dir():
        pytest.skip("Requires sibling parser checkout")
    monkeypatch.syspath_prepend(str(parser_root))
    from parser.company_checks import JobStore, JobWorker
    from parser.serve_export import ExportHandler

    monkeypatch.setattr("parser.serve_export.LAVOK_INGEST_TOKEN", "integration-token")
    monkeypatch.setenv("COMPANY_CHECK_PAUSE", "0")
    server = ThreadingHTTPServer(("127.0.0.1", 0), ExportHandler)
    server.job_store = JobStore(tmp_path)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    monkeypatch.setattr(
        checks.settings,
        "lavok_parser_checks_url",
        f"http://127.0.0.1:{server.server_port}/lavok/checks",
    )
    monkeypatch.setattr(checks.settings, "lavok_parser_ingest_token", "integration-token")
    app = FastAPI()
    app.include_router(checks.router)
    dep = checks.router.routes[0].dependant.dependencies[0].call
    app.dependency_overrides[dep] = lambda: SimpleNamespace(id=42)
    wb = Workbook()
    wb.active.append(["ИНН организации", "Наименование организации"])
    wb.active.append(["7733475417", "Компания Тест"])
    stream = io.BytesIO()
    wb.save(stream)

    def source(source_name, company, job_id):
        if source_name == "buh":
            return {"inn": company["inn"], "years": [{"period": "2025", "revenue": 123000}]}
        if source_name == "kad":
            return {"error": "blocked_451"}
        return {"inn": company["inn"], "ogrn": "1257700417611"}

    try:
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://crm"
        ) as client:
            response = await client.post(
                "/api/v1/lavok-parser/checks", files={"file": ("data.xlsx", stream.getvalue())}
            )
            assert response.status_code == 202
            job_id = response.json()["id"]
            JobWorker(server.job_store, parser_root, source).process(job_id)
            progress = await client.get(f"/api/v1/lavok-parser/checks/{job_id}")
            assert progress.json()["progress"] == 100
            assert progress.json()["status"] == "completed_with_errors"
            result = await client.get(f"/api/v1/lavok-parser/checks/{job_id}/result.xlsx")
            assert result.status_code == 200
            book = load_workbook(io.BytesIO(result.content))
            sheet = book["Компании"]
            values = dict(
                zip([c.value for c in sheet[1]], [c.value for c in sheet[2]], strict=True)
            )
            assert values["ИНН"] == "7733475417"
            assert values["Выручка 2025, ₽"] == 123000
            assert values["КАД: ошибка"] == "blocked_451"
            assert sheet.auto_filter.ref
            book.close()
    finally:
        server.shutdown()
        server.server_close()
