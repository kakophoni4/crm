from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace as S
from unittest.mock import AsyncMock, Mock
import pytest
from openpyxl import Workbook, load_workbook
from app.workers import benik_collection as collector
from app.modules.leads.opt.registry_export import build_registry_workbook


def workbook_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Заявка на НДС'
    ws.append(['ИНН покупателя', 'Стоимость покупки', 'ИНН продавца', 'Дата счета-фактуры'])
    ws.append(['7733418909', 122000, '7733430705', date(2025,11,5)])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_recognizes_beneficiary_by_content():
    parsed = collector.parse_candidate(workbook_bytes())
    assert parsed.buyer_inn == '7733418909'
    assert len(parsed.lines) == 1
    assert parsed.lines[0].document_date == date(2025,11,5)


def test_export_wraps_long_names_without_changing_values():
    name = 'Общество с ограниченной ответственностью Очень длинное наименование организации ' * 3
    order = S(buyer_name=name, buyer_inn='7733418909',buyer_kpp='001234567')
    line = S(document_number='000001',document_date=date(2026,4,1),amount=Decimal('122.01'),vat_amount=Decimal('22.00'),amount_without_vat=Decimal('100.01'),supplier_name=name,supplier_inn='7733430705',supplier_kpp='001234568')
    data = build_registry_workbook(order,[line])
    wb = load_workbook(BytesIO(data))
    ws = wb.active
    assert ws['C2'].value == name
    assert ws['G2'].value == '7733430705'
    assert ws['I2'].value == 122.01
    assert ws['I3'].value == 122.01
    assert ws['C2'].alignment.wrap_text
    assert ws.row_dimensions[2].height > ws.row_dimensions[1].height
    assert collector.parse_candidate(data) is None


@pytest.mark.parametrize('apply',[False,True])
async def test_history_keeps_document_period_and_deal_unchanged(monkeypatch,apply):
    lead = S(id=5,chat_id=4,contact_id=3,custom_fields={'order':{'service':'Деревья'}},closed_at='closed')
    db = AsyncMock()
    db.execute.return_value = Mock(first=Mock(return_value=None))
    db.get.side_effect = [lead,S(telegram_username='client')]
    db.scalar.side_effect = [None,None,9,2]
    storage = S(get_bytes=AsyncMock(return_value=(workbook_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')))
    monkeypatch.setattr(collector,'get_file_storage',lambda:storage)
    repo = S(new_crm_id=Mock(return_value='synthetic-id'),create_order=AsyncMock(return_value=S(id=44)))
    monkeypatch.setattr(collector,'OptOrderRepository',lambda _db:repo)
    result = await collector.collect_attachment(db,S(id=8,lead_id=5,chat_id=4),0,
        dict(filename='request.xlsx',status='ready',storage_key='test-key'),apply=apply,historical=True)
    assert result['period_code'] == '4/25'
    assert lead.custom_fields == {'order':{'service':'Деревья'}}
    assert lead.closed_at == 'closed'
    if apply:
        assert result['status'] == 'imported'
        assert repo.create_order.call_args.kwargs['season_id'] == 2
        assert repo.create_order.call_args.kwargs['order_kind'] == 'benik'
    else:
        assert result['status'] == 'ready_to_import'
        repo.create_order.assert_not_awaited()


async def test_deleted_source_is_not_reimported(monkeypatch):
    db = AsyncMock()
    db.execute.return_value = Mock(first=Mock(return_value=S(id=55,deleted_at='deleted')))
    storage = Mock()
    monkeypatch.setattr(collector,'get_file_storage',storage)
    result = await collector.collect_attachment(db,S(id=8),0,dict(filename='request.xlsx'),apply=True,historical=True)
    assert result['status'] == 'previously_deleted'
    storage.assert_not_called()


def test_rejects_non_spreadsheet_and_oversized_input():
    assert not collector.spreadsheet({'filename':'document.pdf','mime':'application/pdf'})
    with pytest.raises(ValueError,match='excel_unreadable'):
        collector.parse_candidate(b'not excel')
    with pytest.raises(ValueError,match='file_too_large'):
        collector.parse_candidate(b'x' * (20*1024*1024+1))
