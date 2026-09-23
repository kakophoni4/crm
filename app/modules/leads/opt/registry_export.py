from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
import textwrap

from app.modules.db.models.lead_opt_order import LeadOptOrder, LeadOptOrderLine


def build_registry_workbook(order: LeadOptOrder, lines: list[LeadOptOrderLine]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "TDSheet"
    headers = [
        "№ документа",
        "Дата документа",
        "Покупатель",
        "ИНН покупателя",
        "КПП покупателя",
        "Поставщик",
        "ИНН поставщика",
        "КПП поставщика",
        "Сумма",
        "Сумма НДС",
        "Сумма без НДС",
    ]
    worksheet.append(headers)

    total_amount = Decimal("0")
    total_vat = Decimal("0")
    total_wo_vat = Decimal("0")

    for line in lines:
        doc_date = line.document_date
        if isinstance(doc_date, date):
            date_text = doc_date.strftime("%d.%m.%Y")
        else:
            date_text = str(doc_date)
        amount = Decimal(str(line.amount))
        vat = Decimal(str(line.vat_amount))
        wo_vat = Decimal(str(line.amount_without_vat))
        total_amount += amount
        total_vat += vat
        total_wo_vat += wo_vat
        worksheet.append(
            [
                line.document_number or "",
                date_text,
                order.buyer_name or "",
                order.buyer_inn,
                order.buyer_kpp or "",
                line.supplier_name or "",
                line.supplier_inn,
                line.supplier_kpp or "",
                float(amount),
                float(vat),
                float(wo_vat),
            ],
        )

    worksheet.append(
        [
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            float(total_amount),
            float(total_vat),
            float(total_wo_vat),
        ],
    )

    # Excel does not reliably auto-fit wrapped rows on opening generated files.
    # Set widths first, then a conservative height from every displayed cell.
    widths = [20, 18, 48, 18, 18, 48, 18, 18, 22, 22, 22]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row in worksheet.iter_rows():
        line_count = 1
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Calibri", size=11, bold=cell.row == 1)
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if cell.column >= 9 and cell.row > 1:
                cell.number_format = '#,##0.00'
            width = max(1, int(widths[cell.column - 1] / 1.25))
            count = sum(max(1, len(textwrap.wrap(part, width=width))) for part in str(cell.value or '').split('\n'))
            line_count = max(line_count, count)
        worksheet.row_dimensions[row[0].row].height = min(409, 8 + 15 * line_count)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:K{max(1, worksheet.max_row - 1)}"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
