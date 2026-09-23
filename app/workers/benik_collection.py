"""Collect beneficiary applications only; never submit to 1C or alter deal fields.
CLI defaults to read-only preview of messages existing at deployment.
"""
from __future__ import annotations
import argparse
import asyncio
import json
from collections import Counter
from io import BytesIO
from zipfile import ZipFile, is_zipfile
from sqlalchemy import select, text, or_, and_
import structlog
from app.shared.db import get_session_factory, dispose_engine
from app.shared.storage import get_file_storage
from app.shared.exceptions import ValidationError
from app.modules.db.models.chat_message import ChatMessage
from app.modules.db.models.lead import Lead
from app.modules.db.models.lead_opt_order import LeadOptOrder
from app.modules.db.models.user import User
from app.modules.db.models.enums import UserRole, UserStatus
from app.modules.db.models.contact import Contact
from app.modules.db.models.blacklist_entry import BlacklistEntry
from app.modules.leads.opt.repository import OptOrderRepository
from app.modules.leads.opt.nds_request_parser import parse_nds_request_workbook
from app.modules.leads.opt.fingerprint import compute_application_fingerprint
from app.modules.leads.opt.periods import resolve_application_period
from app.modules.leads.opt.vat import split_vat_included, vat_rate_for_period_code

logger = structlog.get_logger(__name__)

def spreadsheet(att):
    name = str(att.get('filename') or att.get('name') or '').lower()
    mime = str(att.get('mime') or '').lower()
    return name.endswith(('.xlsx','.xls')) or 'spreadsheet' in mime or 'ms-excel' in mime

def parse_candidate(content):
    if len(content) > 20 * 1024 * 1024:
        raise ValueError('file_too_large')
    if is_zipfile(BytesIO(content)):
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024 or len(archive.infolist()) > 2000:
                raise ValueError('workbook_too_large')
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(content), read_only=True)
        try:
            if any((s.max_row or 0) > 50000 or (s.max_column or 0) > 200 for s in wb):
                raise ValueError('workbook_dimensions_too_large')
        finally:
            wb.close()
    result = parse_nds_request_workbook(content)
    # Match the existing manual import classification, not just the filename.
    if result.reason and (result.reason.startswith('excel_open_failed') or result.reason == 'xlrd_not_installed'):
        raise ValueError('excel_unreadable')
    if not result.matched or result.form_kind != 'nds_request':
        return None
    if result.application is None:
        raise ValueError(result.reason or 'invalid_application')
    return result.application

async def collect_attachment(db, message, index, att, *, apply, historical):
    result = {'message_id':message.id, 'attachment_index':index, 'filename':str(att.get('filename') or att.get('name') or '')}
    def done(status, **extra):
        return {**result, 'status':status, **extra}
    if not spreadsheet(att): return done('not_excel')
    existing = (await db.execute(select(LeadOptOrder.id, LeadOptOrder.deleted_at).where(
        LeadOptOrder.source_message_id == message.id, LeadOptOrder.source_attachment_index == index
    ).limit(1))).first()
    if existing: return done('previously_deleted' if existing.deleted_at else 'duplicate', order_id=existing.id)
    if att.get('status') != 'ready' or not att.get('storage_key'):
        return done('attachment_unavailable')
    content, _ = await get_file_storage().get_bytes(str(att['storage_key']))
    parsed = await asyncio.to_thread(parse_candidate, content)
    if parsed is None: return done('not_beneficiary')
    fingerprint = compute_application_fingerprint(parsed)
    existing_id = await db.scalar(select(LeadOptOrder.id).where(
        LeadOptOrder.content_fingerprint == fingerprint).order_by(LeadOptOrder.id).limit(1))
    if existing_id: return done('duplicate', order_id=existing_id)
    lead = await db.get(Lead, message.lead_id) if message.lead_id else None
    if lead is None:
        candidates = (await db.scalars(select(Lead).where(Lead.chat_id == message.chat_id).limit(2))).all()
        if len(candidates) != 1: return done('needs_review', reason='ambiguous_or_missing_deal')
        lead = candidates[0]
    if lead.chat_id != message.chat_id: return done('needs_review', reason='deal_chat_mismatch')
    contact = await db.get(Contact, lead.contact_id)
    username = (contact.telegram_username or '').lstrip('@').lower() if contact else ''
    blocked = await db.scalar(select(BlacklistEntry.id).where(or_(
        and_(BlacklistEntry.kind == 'buyer_inn', BlacklistEntry.value == parsed.buyer_inn),
        and_(BlacklistEntry.kind == 'telegram_username', BlacklistEntry.value == username),
    )).limit(1))
    if blocked: return done('blocked')
    period = resolve_application_period([line.document_date for line in parsed.lines]).period_code
    result.update(lead_id=lead.id, period_code=period, line_count=len(parsed.lines))
    if not apply: return done('ready_to_import', season='2-й квартал 2026' if historical else 'current')
    actor_id = await db.scalar(select(User.id).where(User.role == UserRole.ADMIN, User.status == UserStatus.ACTIVE).order_by(User.id).limit(1))
    if actor_id is None: return done('needs_review', reason='no_active_admin')
    season_id = None
    if historical:
        season_id = await db.scalar(text("SELECT id FROM sales_seasons WHERE name='2-й квартал 2026'"))
        if season_id is None: raise ValueError('historical_season_missing')
    # Serialize order numbering and all automatic imports into the same deal.
    await db.execute(select(Lead.id).where(Lead.id == lead.id).with_for_update())
    repo = OptOrderRepository(db)
    rate = vat_rate_for_period_code(period)
    lines = []
    for line in parsed.lines:
        amount, vat, without = split_vat_included(line.amount, rate_percent=rate)
        lines.append(dict(crm_id=repo.new_crm_id('crm-line'), supplier_inn=line.supplier_inn,
            supplier_kpp=line.supplier_kpp, supplier_name=line.supplier_name,
            document_date=line.document_date, amount=float(amount), vat_amount=float(vat), amount_without_vat=float(without)))
    order = await repo.create_order(lead_id=lead.id, crm_id=repo.new_crm_id('crm-order'),
        buyer_inn=parsed.buyer_inn, buyer_kpp=parsed.buyer_kpp, buyer_name=parsed.buyer_name,
        source_filename=result['filename'] or 'application.xlsx', created_by=actor_id,
        lines=lines, source_message_id=message.id, source_attachment_index=index,
        content_fingerprint=fingerprint, vat_rate_percent=float(rate), period_code=period,
        order_kind='benik', season_id=season_id)
    return done('imported', order_id=order.id)

async def collect_message(db, message_id, *, apply, historical):
    message = await db.get(ChatMessage, message_id)
    if message is None or str(message.direction) != 'inbound': return []
    results = []
    for index, att in enumerate(message.attachments or []):
        if not spreadsheet(att): continue
        try:
            async with db.begin_nested():
                result = await collect_attachment(db, message, index, att, apply=apply, historical=historical)
        except (ValueError, ValidationError) as exc:
            result = dict(message_id=message_id, attachment_index=index, status='needs_review', reason=str(exc)[:300])
        except Exception as exc:
            result = dict(message_id=message_id, attachment_index=index, status='retry', reason=type(exc).__name__)
        results.append(result)
    return results

async def run_one():
    async with get_session_factory()() as db:
        row = (await db.execute(text("SELECT message_id,attempts FROM benik_collection_queue WHERE status='pending' AND next_attempt_at<=now() ORDER BY message_id FOR UPDATE SKIP LOCKED LIMIT 1"))).first()
        if row is None: return False
        cutoff = await db.scalar(text('SELECT historical_max_id FROM benik_collection_state WHERE id=1'))
        report = await collect_message(db, row.message_id, apply=True, historical=row.message_id <= cutoff)
        retry = any(r['status'] in ('retry','attachment_unavailable') for r in report)
        state = 'pending' if retry and row.attempts < 11 else 'needs_review' if retry or any(r['status']=='needs_review' for r in report) else 'done'
        await db.execute(text("""UPDATE benik_collection_queue SET status=:state, attempts=attempts+1,
            next_attempt_at=now()+interval '5 minutes', updated_at=now(), report=CAST(:report AS jsonb) WHERE message_id=:id"""),
            dict(state=state,report=json.dumps(report,ensure_ascii=False),id=row.message_id))
        await db.commit()
        return True

async def collection_loop():
    while True:
        try:
            for _ in range(10):
                if not await run_one(): break
        except Exception as exc:
            logger.warning('benik_collection_error', error_type=type(exc).__name__)
        await asyncio.sleep(15)

async def scan_history(apply=False):
    counts = Counter()
    after = 0
    try:
        async with get_session_factory()() as db:
            cutoff = await db.scalar(text('SELECT historical_max_id FROM benik_collection_state WHERE id=1'))
        while True:
            async with get_session_factory()() as db:
                ids = (await db.scalars(select(ChatMessage.id).where(ChatMessage.id > after,
                    ChatMessage.id <= cutoff, ChatMessage.direction == 'inbound',
                    text("jsonb_array_length(attachments)>0")).order_by(ChatMessage.id).limit(100))).all()
            if not ids: break
            for mid in ids:
                async with get_session_factory()() as db:
                    if not apply: await db.execute(text('SET TRANSACTION READ ONLY'))
                    report = await collect_message(db, mid, apply=apply, historical=True)
                    if apply:
                        state = 'needs_review' if any(r['status'] in ('needs_review','retry','attachment_unavailable') for r in report) else 'done'
                        await db.execute(text("""INSERT INTO benik_collection_queue(message_id,status,report) VALUES(:id,:state,CAST(:report AS jsonb))
                            ON CONFLICT(message_id) DO UPDATE SET status=EXCLUDED.status, report=EXCLUDED.report, updated_at=now()"""), dict(id=mid,state=state,report=json.dumps(report,ensure_ascii=False)))
                        await db.commit()
                    for item in report:
                        counts[item['status']] += 1
                        print(json.dumps(item,ensure_ascii=False),flush=True)
            after = ids[-1]
        print(json.dumps({'summary':dict(counts),'applied':apply},ensure_ascii=False),flush=True)
    finally:
        await dispose_engine()

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--apply-history', action='store_true', help='Import historical beneficiary files into season Q2 2026')
    args = parser.parse_args()
    asyncio.run(scan_history(args.apply_history))
